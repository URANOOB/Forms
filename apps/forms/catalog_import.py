"""Read-only CSV/Excel analysis and explicit proposals for dependent choice fields."""

import csv
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from io import BytesIO, StringIO
from itertools import islice
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .catalog_fields import build_fields_proposal, read_definitions, suggest_fields

MAX_OPTIONS = 5000
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_COLUMNS = 50


def normalized(value):
    return "".join(
        ch for ch in unicodedata.normalize("NFD", value.lower()) if unicodedata.category(ch) != "Mn"
    )


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float)) and value == int(value):
        text = str(int(value))
        # Excel commonly stores identifiers as numbers with a zero-padding format.
        if re.fullmatch(r"0{2,150}", cell.number_format):
            return text.zfill(len(cell.number_format))
        return text
    return str(value).strip()


def suggest_columns(headers, rows):
    columns = range(len(headers))
    unique = [len({row[i] for _, row in rows if row[i]}) for i in columns]
    names = [normalized(header) for header in headers]
    code = max(
        columns,
        key=lambda i: (
            bool(re.search(r"\b(codigo|code|id|clave)\b", names[i]))
            and not re.search(r"descripcion|description|nombre", names[i]),
            unique[i],
            -sum(len(row[i]) for _, row in rows),
        ),
    )
    remaining = [i for i in columns if i != code]
    description = (
        max(
            remaining,
            key=lambda i: (
                bool(re.search(r"descripcion|description|nombre|label", names[i])),
                sum(len(row[i]) for _, row in rows),
            ),
        )
        if remaining
        else None
    )
    remaining = [i for i in remaining if i != description]
    group = (
        max(
            remaining,
            key=lambda i: (
                bool(re.search(r"agrupador|grupo|categoria|category|group|familia", names[i])),
                1 < unique[i] < len(rows),
                unique[i],
            ),
        )
        if remaining
        else None
    )
    if (
        len(headers) == 2
        and description is not None
        and re.search(r"agrupador|grupo|categoria|category|group|familia", names[description])
    ):
        group, description = description, None
    return {"code": code, "description": description, "group": group}


def sheet_issue(names, selected, message, header_row=1):
    """Keep sheet/header selection usable even when this sheet has no usable table."""
    return {
        "sheets": names,
        "sheet": selected,
        "header_row": header_row,
        "headers": [],
        "mapping": {"code": None, "description": None, "group": None},
        "row_count": 0,
        "sample": [],
        "issues": [message],
        "issue_count": 1,
        "warnings": [],
        "proposal": None,
        "field_definitions": [],
    }


def analyze_catalog(upload, settings):
    extension = upload.name.lower().rsplit(".", 1)[-1]
    if extension not in {"csv", "xlsx", "xlsm"}:
        raise ValidationError("Selecciona un archivo .csv, .xlsx o .xlsm.")
    if upload.size > MAX_UPLOAD_BYTES:
        raise ValidationError("El archivo admite hasta 5 MB.")
    content = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValidationError("El archivo admite hasta 5 MB.")
    if extension == "csv":
        names, selected, raw = read_csv_rows(content, settings)
    else:
        names, selected, raw = read_excel_rows(content, settings)
    return analyze_rows(names, selected, raw, settings)


def csv_delimiter(text):
    try:
        return csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
    except csv.Error:
        # Sniffer can fail on multiline quoted values or title rows above a table.
        # Compare complete records, respecting quotes, instead of physical lines.
        def score(delimiter):
            try:
                reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
                widths = Counter(len(row) for row in islice(reader, 100) if len(row) > 1)
                return max(widths.values(), default=0)
            except csv.Error:
                return 0

        return max((",", ";", "\t"), key=score)


def read_csv_rows(content, settings):
    names = ["CSV"]
    selected = settings.get("sheet") or names[0]
    if selected not in names:
        raise ValidationError("Selecciona una hoja del archivo.")
    try:
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = content.decode("utf-16")
        else:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")
        if "\x00" in text:
            raise ValueError("Contenido binario")
        raw = []
        reader = csv.reader(StringIO(text, newline=""), delimiter=csv_delimiter(text), strict=True)
        for index, row in enumerate(reader, 1):
            if index > MAX_OPTIONS + 50:
                raise ValidationError(
                    f"La hoja admite hasta {MAX_OPTIONS} filas y 50 de encabezado."
                )
            values = [value.strip() for value in row]
            if any(values[MAX_COLUMNS:]):
                raise ValidationError(f"La tabla admite hasta {MAX_COLUMNS} columnas.")
            values = (values[:MAX_COLUMNS] + [""] * MAX_COLUMNS)[:MAX_COLUMNS]
            raw.append((index, values, set()))
        return names, selected, raw
    except (UnicodeError, ValueError, csv.Error) as error:
        raise ValidationError(
            "No se pudo leer el CSV. Revisa su codificación y formato."
        ) from error


def read_excel_rows(content, settings):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
                raise ValidationError("El contenido del Excel es demasiado grande.")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, KeyError, ValueError, OSError) as error:
        raise ValidationError(
            "No se pudo leer el Excel. Revisa que sea un .xlsx o .xlsm válido."
        ) from error
    try:
        names = workbook.sheetnames
        if not names:
            raise ValidationError("El archivo no contiene hojas.")
        if len(names) > 20:
            raise ValidationError("El archivo admite hasta 20 hojas.")
        selected = settings.get("sheet") or names[0]
        if selected not in names:
            raise ValidationError("Selecciona una hoja del archivo.")
        sheet = workbook[selected]
        # Ignore inflated/incorrect dimension metadata; bound actual iteration ourselves.
        sheet.reset_dimensions()
        raw = []
        for index, cells in enumerate(sheet.iter_rows(), 1):
            if index > MAX_OPTIONS + 50:
                raise ValidationError(
                    f"La hoja admite hasta {MAX_OPTIONS} filas y 50 de encabezado."
                )
            values = [cell_text(cell) for cell in cells]
            if any(values[MAX_COLUMNS:]):
                raise ValidationError(f"La tabla admite hasta {MAX_COLUMNS} columnas.")
            invalid = {i for i, cell in enumerate(cells) if cell.data_type in {"f", "e"}}
            values = (values[:MAX_COLUMNS] + [""] * MAX_COLUMNS)[:MAX_COLUMNS]
            raw.append(
                (
                    index,
                    values,
                    invalid,
                )
            )
        return names, selected, raw
    finally:
        workbook.close()


def analyze_rows(names, selected, raw, settings):
    populated = [(n, row, errors) for n, row, errors in raw if any(row)]
    if not populated:
        return sheet_issue(names, selected, "La hoja está vacía. Selecciona otra hoja.")
    header_row = settings.get("header_row")
    if header_row in (None, ""):
        candidates = [item for item in populated if item[0] <= 50]
        if not candidates:
            raise ValidationError("Los encabezados deben estar en las primeras 50 filas.")
        header_row = max(
            candidates,
            key=lambda item: (
                sum(
                    bool(re.search(r"codigo|descripcion|agrupador|categoria|nombre", normalized(v)))
                    for v in item[1]
                ),
                sum(bool(v) for v in item[1]),
                -item[0],
            ),
        )[0]
    try:
        header_row = int(header_row)
    except (TypeError, ValueError) as error:
        raise ValidationError("Indica la fila de encabezados (1 a 50).") from error
    if not 1 <= header_row <= min(50, len(raw)):
        raise ValidationError("Indica una fila de encabezados existente entre 1 y 50.")
    header = raw[header_row - 1][1]
    width = max((i + 1 for i, value in enumerate(header) if value), default=0)
    if width < 1:
        return sheet_issue(names, selected, "Selecciona una fila con encabezados.", header_row)
    headers = [value or f"Columna {i + 1}" for i, value in enumerate(header[:width])]
    records = [(n, row[:width]) for n, row, _ in populated if n > header_row]
    if not records:
        return sheet_issue(
            names, selected, "No hay filas de datos bajo los encabezados.", header_row
        )
    if len(records) > MAX_OPTIONS:
        raise ValidationError(f"Selecciona una tabla de 1 a {MAX_OPTIONS} filas de datos.")
    mapping = suggest_columns(headers, records)
    if "code" in settings:
        try:
            mapping = {
                key: int(settings[key]) if settings.get(key) != "" else None
                for key in ("code", "description", "group")
            }
        except (TypeError, ValueError, KeyError) as error:
            raise ValidationError("Selecciona las columnas para crear los campos.") from error
    if "code" in settings:
        indexes = [i for i in mapping.values() if i is not None]
        if any(i < 0 or i >= width for i in indexes) or len(set(indexes)) != len(indexes):
            raise ValidationError("Usa columnas distintas para código, descripción y agrupador.")
        if mapping["code"] is None:
            raise ValidationError("Selecciona la columna del valor guardado.")
    defaults = suggest_fields(headers, mapping)
    # Accept the previous mapping request format as well as explicit field definitions.
    if "code" in settings:
        for field in defaults:
            prefix = "group" if field["value_column"] == mapping["group"] else "item"
            field["label"] = settings.get(f"{prefix}_label", "").strip() or field["label"]
            field["searchable"] = settings.get(f"{prefix}_searchable", "true") == "true"
    definitions = read_definitions(settings, headers, defaults)
    indexes = {
        column
        for field in definitions
        for column in [field["value_column"], *field["label_columns"]]
    }
    issues = []
    for n, _, invalid in populated:
        if n > header_row and invalid.intersection(indexes):
            issues.append(
                f"Fila {n}: hay una fórmula o un error en las columnas elegidas. Usa valores."
            )
    proposal = None
    warnings = []
    if not issues:
        proposal, issues, warnings = build_fields_proposal(records, definitions)
    return {
        "sheets": names,
        "sheet": selected,
        "header_row": header_row,
        "headers": headers,
        "mapping": mapping,
        "field_definitions": definitions,
        "row_count": len(records),
        "sample": [row for _, row in records[:5]],
        "issues": issues[:20],
        "issue_count": len(issues),
        "warnings": warnings,
        "proposal": proposal,
    }
