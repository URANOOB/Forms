import copy
import csv
import io
import json
from zipfile import ZipFile

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.xml.constants import XLSM

from apps.accounts.models import User, Workspace
from apps.forms.builder import current_version, document, save_document
from apps.forms.catalog_import import analyze_catalog
from apps.forms.conditions import FormSchema
from apps.forms.models import Form, FormVersion
from apps.forms.publication import publish_form
from apps.submissions.runtime import PublicResponseForm


def excel(rows=None, prepare=None, name="catalogo.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catálogo"
    for row in rows or [
        ["Código", "Descripción", "Agrupador"],
        ["001", "Primero", "Grupo A"],
        ["002", "Segundo", "Grupo A"],
        ["003", "Tercero", "Grupo B"],
    ]:
        sheet.append(row)
    if prepare:
        prepare(workbook)
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(name, output.getvalue())


def csv_file(rows, delimiter=",", encoding="utf-8-sig"):
    output = io.StringIO(newline="")
    csv.writer(output, delimiter=delimiter).writerows(rows)
    return SimpleUploadedFile("catalogo.csv", output.getvalue().encode(encoding))


def definition(key, column, parent="", labels=None, **kwargs):
    return {
        "id": key,
        "label": key.title(),
        "value_column": column,
        "label_columns": labels if labels is not None else [column],
        "parent": parent,
        "searchable": True,
        "required": False,
        **kwargs,
    }


def configured(rows, definitions):
    return analyze_catalog(excel(rows), {"fields": json.dumps(definitions)})


class CatalogAnalysisTests(SimpleTestCase):
    def test_csv_dialects_and_encodings_preserve_identifiers_and_quoted_labels(self):
        rows = [["Código", "Descripción"], ["001", "Niñez, salud; atención\nIntegral"]]
        for delimiter in (",", ";", "\t"):
            for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1252"):
                with self.subTest(delimiter=delimiter, encoding=encoding):
                    result = analyze_catalog(csv_file(rows, delimiter, encoding), {})
                    self.assertEqual(result["headers"], rows[0])
                    self.assertEqual(result["sample"], rows[1:])
                    self.assertEqual(result["issues"], [])
                    self.assertEqual(result["proposal"]["fields"][0]["options"][0]["value"], "001")

    def test_csv_single_column_empty_and_custom_headers(self):
        result = analyze_catalog(csv_file([["Color"], ["Rojo"], ["Azul"]]), {})
        self.assertEqual(result["proposal"]["field_count"], 1)
        self.assertEqual(result["row_count"], 2)
        result = analyze_catalog(csv_file([]), {})
        self.assertIsNone(result["proposal"])
        result = analyze_catalog(
            csv_file([["Título"], ["Código", "Nombre"], ["001", "Uno"]], ";"),
            {"header_row": "2"},
        )
        self.assertEqual(result["headers"], ["Código", "Nombre"])
        self.assertEqual(result["row_count"], 1)

    def test_csv_limits_and_malformed_files(self):
        for upload in (
            csv_file([["Código"], *[[str(i)] for i in range(5001)]]),
            csv_file([[f"Columna {i}" for i in range(51)], ["dato"] * 51]),
            SimpleUploadedFile("bad.csv", b'Codigo,Nombre\n001,"unterminated'),
            SimpleUploadedFile("bad.csv", b"\x00\x01\x02"),
            SimpleUploadedFile("bad.txt", b"Codigo\n001"),
        ):
            with self.subTest(name=upload.name), self.assertRaises(ValidationError):
                analyze_catalog(upload, {})
        for extension in ("csv", "xlsx", "xlsm"):
            with (
                self.subTest(extension=extension),
                self.assertRaisesMessage(ValidationError, "hasta 5 MB"),
            ):
                upload = SimpleUploadedFile(f"large.{extension}", b"x" * (5 * 1024 * 1024 + 1))
                analyze_catalog(upload, {})

    def test_xlsm_reads_sheets_and_validates_formulas(self):
        # Match a macro-enabled workbook's content type without executing any VBA.
        source = excel()
        output = io.BytesIO()
        with ZipFile(io.BytesIO(source.read())) as original, ZipFile(output, "w") as archive:
            for entry in original.infolist():
                content = original.read(entry.filename)
                if entry.filename == "[Content_Types].xml":
                    content = content.replace(
                        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                        XLSM.encode(),
                    )
                archive.writestr(entry, content)
        result = analyze_catalog(SimpleUploadedFile("catalogo.XLSM", output.getvalue()), {})
        self.assertEqual(result["proposal"]["option_count"], 5)
        self.assertEqual(result["sheets"], ["Catálogo"])
        result = analyze_catalog(
            excel([["Código", "Descripción"], ["=1+1", "Dos"]], name="catalogo.xlsm"), {}
        )
        self.assertIsNone(result["proposal"])
        self.assertIn("fórmula", result["issues"][0])

    def test_single_column_becomes_one_independent_field(self):
        result = analyze_catalog(excel([["Color"], ["Rojo"], ["Azul"], ["Rojo"]]), {})
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["proposal"]["field_count"], 1)
        field = result["proposal"]["fields"][0]
        self.assertEqual([option["value"] for option in field["options"]], ["Rojo", "Azul"])
        self.assertNotIn("option_filter", field["configuration"])
        self.assertTrue(result["warnings"])

    def test_two_columns_suggest_one_combined_or_two_related_fields(self):
        for header, expected in [("Descripción", 1), ("Grupo", 2)]:
            with self.subTest(header=header):
                result = analyze_catalog(excel([["Código", header], ["001", "Uno"]]), {})
                self.assertEqual(result["proposal"]["field_count"], expected)

    def test_columns_can_make_independent_fields_and_ordered_combined_labels(self):
        result = configured(
            [["Código", "Descripción", "Grupo"], ["001", "Primero", "A"]],
            [
                definition("opcion", 0, labels=[1, 0, 2], required=True, searchable=False),
                definition("grupo", 2),
            ],
        )
        first, second = result["proposal"]["fields"]
        self.assertEqual(first["options"][0]["label"], "Primero — 001 — A")
        self.assertEqual(first["options"][0]["value"], "001")
        self.assertTrue(first["required"])
        self.assertFalse(first["configuration"]["searchable"])
        self.assertFalse(second["required"])
        self.assertNotIn("option_filter", second["configuration"])

    def test_only_selected_columns_are_validated(self):
        rows = [["Código", "Descripción", "Ignorada"], ["001", "Primero", "=1+1"]]
        self.assertIsNotNone(configured(rows, [definition("codigo", 0)])["proposal"])
        result = configured(rows, [definition("codigo", 0, labels=[0, 2])])
        self.assertIsNone(result["proposal"])
        self.assertIn("Fila 2", result["issues"][0])

    def test_removed_parent_can_leave_one_independent_field(self):
        result = configured(None, [definition("codigo", 0, labels=[0, 1])])
        field = result["proposal"]["fields"][0]
        self.assertEqual(len(field["options"]), 3)
        self.assertNotIn("option_filter", field["configuration"])

    def test_intermediate_duplicate_values_cannot_leak_descendants(self):
        rows = [
            ["País", "Ciudad", "Barrio"],
            ["A", "Centro", "Norte"],
            ["B", "Centro", "Sur"],
            ["A", "Otra", "Este"],
        ]
        result = configured(
            rows,
            [
                definition("pais", 0),
                definition("ciudad", 1, "pais"),
                definition("barrio", 2, "ciudad"),
            ],
        )
        country, city, neighborhood = result["proposal"]["fields"]
        self.assertEqual(len(city["options"]), 3)
        city_values = city["configuration"]["option_filter"]["values"]
        children = neighborhood["configuration"]["option_filter"]["values"]
        north, south = children["Norte"][0], children["Sur"][0]
        self.assertNotEqual(north, south)
        self.assertEqual(city_values[north], ["A"])
        self.assertEqual(city_values[south], ["B"])
        self.assertEqual(city_values["Otra"], ["A"])
        self.assertEqual(country["options"][0]["value"], "A")
        self.assertIn("identificadores internos", result["warnings"][0])

    def test_branching_dependencies_and_four_levels(self):
        definitions = [
            definition("a", 0),
            definition("b", 1, "a"),
            definition("c", 2, "b"),
            definition("d", 3, "c"),
            definition("e", 3, "a"),
        ]
        result = configured([["A", "B", "C", "D"], ["1", "2", "3", "4"]], definitions)
        self.assertEqual(result["issues"], [])
        fields = result["proposal"]["fields"]
        self.assertEqual(
            fields[4]["configuration"]["option_filter"]["source"], fields[0]["stable_key"]
        )
        self.assertEqual(
            fields[3]["configuration"]["option_filter"]["source"], fields[2]["stable_key"]
        )

    def test_rejects_invalid_definitions_cycles_forward_links_and_limits(self):
        invalid = [
            [],
            {},
            [None],
            [definition("a", 0)] * 21,
            [definition("a", 0), definition("a", 1)],
            [definition("a", 0, "a")],
            [definition("a", 0, "b"), definition("b", 1)],
            [definition("a", 0, "missing")],
            [definition("a", 99)],
            [definition("a", True)],
            [definition("a", 0, labels=[])],
            [definition("a", 0, labels=[0, 0])],
            [definition("a", 0, labels=[{}])],
            [definition("a", 0, searchable="true")],
            [definition("a", 0, label=" ")],
        ]
        for definitions in invalid:
            with self.subTest(definitions=definitions), self.assertRaises(ValidationError):
                configured(None, definitions)
        with self.assertRaises(ValidationError):
            analyze_catalog(excel(), {"fields": "bad json"})
        result = configured(None, [definition(str(i), 0) for i in range(20)])
        self.assertEqual(result["proposal"]["field_count"], 20)

    def test_suggests_code_label_and_group_and_keeps_identifiers(self):
        result = analyze_catalog(excel(), {})
        self.assertEqual(result["mapping"], {"code": 0, "description": 1, "group": 2})
        self.assertEqual(result["issues"], [])
        proposal = result["proposal"]
        self.assertEqual((proposal["field_count"], proposal["option_count"]), (2, 5))
        parent, child = proposal["fields"]
        self.assertEqual(child["options"][0]["value"], "001")
        self.assertEqual(child["options"][0]["label"], "001 — Primero")
        self.assertEqual(child["configuration"]["option_filter"]["source"], parent["stable_key"])

    def test_uses_selected_sheet_header_and_mapping(self):
        def prepare(workbook):
            other = workbook.create_sheet("Otra")
            other.append(["Título"])
            other.append(["Familia", "Nombre", "Clave"])
            other.append(["Uno", "Artículo", "A01"])

        result = analyze_catalog(
            excel(prepare=prepare),
            {
                "sheet": "Otra",
                "header_row": "2",
                "code": "2",
                "description": "1",
                "group": "0",
                "group_label": "Familia",
                "item_label": "Artículo",
                "required": "true",
                "group_searchable": "false",
            },
        )
        self.assertEqual(result["row_count"], 1)
        parent, child = result["proposal"]["fields"]
        self.assertTrue(child["required"])
        self.assertFalse(parent["configuration"]["searchable"])
        self.assertEqual(child["label"], "Artículo")
        self.assertEqual(child["options"][0]["value"], "A01")

    def test_preserves_numeric_identifier_with_zero_format(self):
        def prepare(workbook):
            workbook.active["A2"] = 12
            workbook.active["A2"].number_format = "00000"

        child = analyze_catalog(excel(prepare=prepare), {})["proposal"]["fields"][1]
        self.assertEqual(child["options"][0]["value"], "00012")

    def test_duplicates_are_reported_and_many_to_many_is_preserved(self):
        result = analyze_catalog(
            excel(
                [
                    ["Código", "Descripción", "Agrupador"],
                    ["A", "Artículo", "Uno"],
                    ["A", "Artículo", "Uno"],
                    ["A", "Artículo", "Dos"],
                ]
            ),
            {},
        )
        self.assertEqual(result["proposal"]["option_count"], 3)
        self.assertEqual(len(result["warnings"]), 2)
        self.assertEqual(
            result["proposal"]["fields"][1]["configuration"]["option_filter"]["values"],
            {"A": ["Uno", "Dos"]},
        )

    def test_issues_block_proposal_without_dropping_bad_rows(self):
        for row in (
            ["001", "Conflicto", "Grupo A"],
            ["004", "", "Grupo A"],
            ["=1+1", "Fórmula", "Grupo A"],
            ["X" * 151, "Largo", "Grupo A"],
        ):
            with self.subTest(row=row):
                result = analyze_catalog(excel(prepare=lambda w: w.active.append(row)), {})
                self.assertIsNone(result["proposal"])
                self.assertGreater(result["issue_count"], 0)

    def test_two_columns_can_be_mapped_without_description(self):
        result = analyze_catalog(
            excel([["Código", "Grupo"], ["01", "Uno"]]),
            {"code": "0", "description": "", "group": "1"},
        )
        self.assertEqual(result["proposal"]["fields"][1]["options"][0]["label"], "01")

    def test_rejects_invalid_mapping_and_files(self):
        for settings in (
            {"sheet": "Falta"},
            {"header_row": "0"},
            {"code": "0", "description": "1", "group": "0"},
            {"code": "99", "description": "1", "group": "2"},
        ):
            with self.subTest(settings=settings), self.assertRaises(ValidationError):
                analyze_catalog(excel(), settings)
        with self.assertRaises(ValidationError):
            analyze_catalog(SimpleUploadedFile("archivo.xlsx", b"bad zip"), {})

    def test_empty_sheet_keeps_sheet_selection_available(self):
        def prepare(workbook):
            workbook.create_sheet("Vacía", 0)

        result = analyze_catalog(excel(prepare=prepare), {})
        self.assertEqual(result["sheets"], ["Vacía", "Catálogo"])
        self.assertIsNone(result["proposal"])
        self.assertTrue(result["issues"])

    def test_rejects_wide_tables_without_silently_truncating_columns(self):
        def prepare(workbook):
            workbook.active.cell(1, 60, "Agrupador distante")

        with self.assertRaises(ValidationError):
            analyze_catalog(excel(prepare=prepare), {})

    def test_row_limit_is_enforced(self):
        def prepare(workbook):
            for i in range(5000):
                workbook.active.append([str(i), "Artículo", "Grupo"])

        with self.assertRaises(ValidationError):
            analyze_catalog(excel(prepare=prepare), {})


class CatalogIntegrationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.workspace = Workspace.objects.create(name="Equipo", slug="catalogos")
        cls.user = User.objects.create_user(
            username="catalog-editor", is_staff=True, workspace=cls.workspace
        )
        cls.user.user_permissions.set(
            Permission.objects.filter(
                content_type__app_label="forms",
                codename__in=["change_form", "add_form", "view_form"],
            )
        )
        cls.form = Form.objects.create(
            workspace=cls.workspace, created_by=cls.user, name="Catálogo", slug="catalogo"
        )
        cls.version = FormVersion.objects.create(form=cls.form, version_number=1)

    def setUp(self):
        self.client.force_login(self.user)
        self.fields = analyze_catalog(excel(), {"required": "true"})["proposal"]["fields"]

    def save_catalog(self, fields=None, rules=None):
        data = document(current_version(self.form))
        data["sections"] = [
            {
                "id": "section",
                "title": "Catálogo",
                "description": "",
                "configuration": {},
                "fields": fields or self.fields,
            }
        ]
        data["rules"] = rules or []
        return save_document(self.form.pk, data)

    def test_configured_chain_survives_save_and_validates_each_level(self):
        definitions = [
            definition("pais", 0, required=True),
            definition("ciudad", 1, "pais", required=True),
            definition("barrio", 2, "ciudad", required=True),
        ]
        result = configured(
            [["País", "Ciudad", "Barrio"], ["A", "Centro", "Norte"], ["B", "Centro", "Sur"]],
            definitions,
        )
        saved = self.save_catalog(result["proposal"]["fields"])
        save_document(self.form.pk, saved)
        schema = FormSchema(current_version(self.form))
        country, city, neighborhood = schema.fields
        mapping = schema.option_filters[str(neighborhood.pk)]["values"]
        north_city, south_city = mapping["Norte"][0], mapping["Sur"][0]
        for country_value, city_value, neighborhood_value, valid in [
            ("A", north_city, "Norte", True),
            ("B", south_city, "Sur", True),
            ("A", north_city, "Sur", False),
            ("B", north_city, "Norte", False),
            ("A", "", "Norte", False),
            ("", "", "", False),
        ]:
            with self.subTest(values=(country_value, city_value, neighborhood_value)):
                response = PublicResponseForm(
                    schema,
                    data={
                        f"answer_{country.stable_key}": country_value,
                        f"answer_{city.stable_key}": city_value,
                        f"answer_{neighborhood.stable_key}": neighborhood_value,
                    },
                )
                self.assertEqual(response.is_valid(), valid, response.errors)

    def test_single_independent_catalog_can_save_and_submit(self):
        result = analyze_catalog(excel([["Color"], ["Rojo"], ["Azul"]]), {})
        self.save_catalog(result["proposal"]["fields"])
        schema = FormSchema(current_version(self.form))
        field = schema.fields[0]
        response = PublicResponseForm(schema, data={f"answer_{field.stable_key}": "Azul"})
        self.assertTrue(response.is_valid(), response.errors)
        self.assertEqual(schema.option_filters, {})

    def test_analysis_endpoint_is_read_only_and_requires_permission_and_csrf(self):
        url = reverse("admin:forms_form_import_catalog")
        before = Form.objects.count(), FormVersion.objects.count()
        response = self.client.post(url, {"file": excel()})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["proposal"]["option_count"], 5)
        self.assertEqual(before, (Form.objects.count(), FormVersion.objects.count()))
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.user)
        self.assertEqual(strict.post(url, {"file": excel()}).status_code, 403)
        self.user.user_permissions.clear()
        self.assertEqual(self.client.post(url, {"file": excel()}).status_code, 403)

    def test_analysis_endpoint_accepts_csv_and_xlsm_and_recovers_after_bad_upload(self):
        url = reverse("admin:forms_form_import_catalog")
        response = self.client.post(url, {"file": SimpleUploadedFile("bad.xlsm", b"invalid")})
        self.assertEqual(response.status_code, 400)
        for upload in (csv_file([["Código"], ["001"]]), excel(name="catalogo.xlsm")):
            with self.subTest(name=upload.name):
                response = self.client.post(url, {"file": upload})
                self.assertEqual(response.status_code, 200)
                self.assertIsNotNone(response.json()["proposal"])

    def test_saves_more_than_one_hundred_options_and_keeps_relation_after_resave(self):
        rows = [
            ["Código", "Descripción", "Agrupador"],
            *[[f"{i:04}", f"Artículo {i}", f"Grupo {i % 3}"] for i in range(125)],
        ]
        self.fields = analyze_catalog(excel(rows), {})["proposal"]["fields"]
        first = self.save_catalog()
        second = save_document(self.form.pk, first)
        self.assertNotEqual(
            first["sections"][0]["fields"][0]["id"], second["sections"][0]["fields"][0]["id"]
        )
        schema = FormSchema(current_version(self.form))
        child = schema.fields[1]
        self.assertEqual(child.options.count(), 125)
        self.assertEqual(schema.option_filters[str(child.pk)]["source"], str(schema.fields[0].pk))

    def test_response_accepts_only_matching_group_and_code(self):
        self.save_catalog()
        schema = FormSchema(current_version(self.form))
        parent, child = schema.fields
        for group, code, expected in [
            ("Grupo A", "001", True),
            ("Grupo A", "003", False),
            ("Grupo B", "003", True),
            ("", "001", False),
            ("Grupo A", "", False),
            ("bad", "001", False),
        ]:
            with self.subTest(group=group, code=code):
                response = PublicResponseForm(
                    schema,
                    data={
                        f"answer_{parent.stable_key}": group,
                        f"answer_{child.stable_key}": code,
                    },
                )
                self.assertEqual(response.is_valid(), expected, response.errors)
        self.assertIn("option_filter", schema.browser_spec()["fields"][str(child.pk)])

    def test_invalid_dependent_answer_cannot_activate_downstream_condition(self):
        target = {
            **copy.deepcopy(self.fields[0]),
            "id": "target",
            "stable_key": "target",
            "field_type": "SHORT_TEXT",
            "configuration": {},
            "options": [],
            "required": True,
        }
        self.save_catalog(
            [*self.fields, target],
            [
                {
                    "source": self.fields[1]["id"],
                    "target_field": "target",
                    "action": "SHOW",
                    "operator": "EQUALS",
                    "expected": "003",
                }
            ],
        )
        schema = FormSchema(current_version(self.form))
        parent, child, target = schema.fields
        states = schema.states({str(parent.pk): "Grupo A", str(child.pk): "003"})
        self.assertFalse(states[str(target.pk)]["visible"])
        states = schema.states({str(parent.pk): "Grupo B", str(child.pk): "003"})
        self.assertTrue(states[str(target.pk)]["visible"])

    def test_rejects_broken_relations_atomically(self):
        original = self.save_catalog()
        for mutation in ("missing_source", "bad_group", "missing_code", "wrong_type", "order"):
            data = copy.deepcopy(original)
            fields = data["sections"][0]["fields"]
            config = fields[1]["configuration"]["option_filter"]
            if mutation == "missing_source":
                config["source"] = "missing"
            if mutation == "bad_group":
                config["values"]["001"] = ["missing"]
            if mutation == "missing_code":
                del config["values"]["001"]
            if mutation == "wrong_type":
                fields[1]["configuration"]["widget"] = "radio"
            if mutation == "order":
                fields.reverse()
            with self.subTest(mutation=mutation), self.assertRaises(ValidationError):
                save_document(self.form.pk, data)
            self.assertEqual(
                document(current_version(self.form))["fingerprint"], original["fingerprint"]
            )

    def test_published_catalog_is_unchanged_after_edit(self):
        self.save_catalog()
        published = publish_form(self.form.pk)
        old_version = published.active_version
        data = document(old_version)
        data["sections"][0]["fields"][1]["options"][0]["label"] = "001 — Nombre nuevo"
        save_document(self.form.pk, data)
        old_child = old_version.fields.get(stable_key=self.fields[1]["stable_key"])
        self.assertEqual(old_child.options.get(value="001").label, "001 — Primero")
        published.refresh_from_db()
        self.assertNotEqual(published.active_version_id, old_version.pk)
        FormSchema(published.active_version)

    def test_hidden_parent_cannot_supply_a_dependent_answer(self):
        toggle = {
            **copy.deepcopy(self.fields[0]),
            "id": "toggle",
            "stable_key": "toggle",
            "label": "Activar",
            "required": False,
            "configuration": {"widget": "select"},
        }
        self.save_catalog(
            [toggle, *self.fields],
            [
                {
                    "source": "toggle",
                    "target_field": self.fields[0]["id"],
                    "action": "SHOW",
                    "operator": "EQUALS",
                    "expected": "Grupo A",
                }
            ],
        )
        schema = FormSchema(current_version(self.form))
        _, parent, child = schema.fields
        response = PublicResponseForm(
            schema,
            data={
                "answer_toggle": "Grupo B",
                f"answer_{parent.stable_key}": "Grupo A",
                f"answer_{child.stable_key}": "001",
            },
        )
        self.assertFalse(response.is_valid())
        self.assertIn(f"answer_{child.stable_key}", response.errors)

    def test_bulk_options_reject_duplicates_and_invalid_flags(self):
        original = self.save_catalog()
        for mutation in ("duplicate", "empty_label", "flag"):
            data = copy.deepcopy(original)
            options = data["sections"][0]["fields"][1]["options"]
            if mutation == "duplicate":
                options[1]["value"] = options[0]["value"]
            elif mutation == "empty_label":
                options[0]["label"] = " "
            else:
                options[0]["is_active"] = "yes"
            with self.subTest(mutation=mutation), self.assertRaises(ValidationError):
                save_document(self.form.pk, data)
            self.assertEqual(
                document(current_version(self.form))["fingerprint"], original["fingerprint"]
            )
