"""Date-scoped Excel reports and private document bundles for staff."""

import hashlib
import tempfile
import uuid
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from itertools import islice
from zipfile import ZIP_DEFLATED, ZipFile

from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.http import (
    FileResponse,
    HttpResponseBadRequest,
    HttpResponseNotAllowed,
    StreamingHttpResponse,
)
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.accounts.models import ReportDownload
from apps.forms.response_summary import SUMMARY_TYPES

from .models import SubmissionAnswer, SubmissionFile
from .summary import answer_text, load_summary_data, normalized, role


def safe_cell(value):
    """Keep submitted text as text, including strings that look like formulas."""
    value = "" if value is None else str(value)
    # XML 1.0 cannot represent these characters; preserve their identity visibly.
    value = ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\u{ord(match[0]):04x}", value)
    value = value.replace("\ufffe", "\\ufffe").replace("\uffff", "\\uffff")
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def report_answer_text(answer):
    value = answer.value
    if (
        answer.field.field_type == "NUMBER"
        and isinstance(value, (int, float))
        and not isinstance(value, bool)
    ):
        try:
            number = Decimal(str(value))
            if number.is_finite() and number == number.to_integral_value():
                return str(int(number))
        except (InvalidOperation, ValueError, OverflowError):
            pass
    return answer_text(answer)


def identity(submission):
    names = []
    surnames = []
    document = ""
    for answer in submission.summary_answers:
        if answer.field.field_type not in SUMMARY_TYPES:
            continue
        value = report_answer_text(answer).strip()
        kind = role(answer)
        if kind == "name" and value:
            names.append(value)
        elif kind == "surname" and value:
            surnames.append(value)
        elif kind == "document" and value and not document:
            document = value
    return " ".join([*names, *surnames]) or "Sin nombre", document


def date_range(request):
    values = []
    for key in ("desde", "hasta"):
        raw = request.GET.get(key, "").strip()
        try:
            values.append(date.fromisoformat(raw) if raw else None)
        except ValueError:
            return None
    start, end = values
    if start and end and start > end:
        return None
    return start, end


REPORT_ORDER = {
    "newest": ("-submitted_at", "-pk"),
    "oldest": ("submitted_at", "pk"),
    "form": ("form__name", "-submitted_at", "-pk"),
}

REPORT_SHEETS = {
    "respuestas": [
        "Recibida (Bogotá)",
    ],
    "documentos": [
        "Nombre",
        "Documento",
        "Campo",
        "Archivo original",
        "Tamaño (bytes)",
        "Ruta en ZIP",
    ],
}


def report_filters(request):
    period = date_range(request)
    if period is None:
        return None
    raw_form = request.GET.get("form", "").strip()
    try:
        form_id = uuid.UUID(raw_form) if raw_form else None
    except ValueError:
        return None
    order = request.GET.get("orden", "newest")
    if order not in REPORT_ORDER:
        return None
    return period, form_id, order


def scoped_queryset(model_admin, request, filters):
    if not model_admin.has_view_permission(request):
        raise PermissionDenied
    period, form_id, order = filters
    start, end = period
    queryset = model_admin.get_queryset(request).select_related("form")
    if form_id:
        queryset = queryset.filter(form_id=form_id)
    tz = timezone.get_current_timezone()
    if start:
        queryset = queryset.filter(
            submitted_at__gte=timezone.make_aware(datetime.combine(start, time.min), tz)
        )
    if end and end < date.max:
        queryset = queryset.filter(
            submitted_at__lt=timezone.make_aware(
                datetime.combine(end + timedelta(days=1), time.min), tz
            )
        )
    return queryset.order_by(*REPORT_ORDER[order])


def scoped_submissions(model_admin, request, filters):
    return prepared_submissions(
        model_admin, request, scoped_queryset(model_admin, request, filters)
    )


def prepared_submissions(model_admin, request, queryset):
    records = queryset.iterator(chunk_size=200)
    while batch := list(islice(records, 200)):
        for submission in batch:
            if not model_admin.has_view_permission(request, submission):
                raise PermissionDenied
        load_summary_data(batch)
        yield from batch


def safe_part(value, fallback):
    return slugify(value, allow_unicode=True)[:70].strip("-._") or fallback


def archive_path(submission, answer, attachment):
    name, document = identity(submission)
    submitted = timezone.localtime(submission.submitted_at).strftime("%Y%m%d-%H%M%S")
    response_suffix = hashlib.sha256(submission.pk.bytes).hexdigest()[:12]
    folder = "_".join(
        (
            safe_part(name, "sin-nombre"),
            safe_part(document, "sin-documento"),
            submitted,
            response_suffix,
        )
    )
    original = attachment.original_name.replace("\\", "/").split("/")[-1]
    field = safe_part(answer.field.label, "archivo")
    filename = safe_part(original.rsplit(".", 1)[0], "archivo")
    extension = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    extension = "".join(
        character for character in extension if character.isascii() and character.isalnum()
    )[:10]
    suffix = f".{extension}" if extension else ""
    file_suffix = hashlib.sha256(attachment.pk.bytes).hexdigest()[:12]
    return f"{folder}/{field}_{filename}_{file_suffix}{suffix}"


def files_for(submissions):
    for submission in submissions:
        for answer in submission.summary_answers:
            for attachment in answer.files.all():
                yield submission, answer, attachment, archive_path(submission, answer, attachment)


class ZipSink:
    """Let zipfile write to an HTTP iterator without holding the archive in memory."""

    def __init__(self):
        self.chunks = []

    def write(self, data):
        self.chunks.append(data)
        return len(data)

    def flush(self):
        pass

    def drain(self):
        chunks, self.chunks = self.chunks, []
        return chunks


def zip_chunks(submissions):
    sink = ZipSink()
    with ZipFile(sink, "w", compression=ZIP_DEFLATED, allowZip64=True) as archive:
        for _, _, attachment, path in files_for(submissions):
            with (
                attachment.file.open("rb") as source,
                archive.open(path, "w", force_zip64=True) as target,
            ):
                while chunk := source.read(64 * 1024):
                    target.write(chunk)
                    yield from sink.drain()
            yield from sink.drain()
    yield from sink.drain()


def answer_key(label):
    return normalized(label) or "campo sin nombre"


def report_columns(queryset):
    answers = (
        SubmissionAnswer.objects.filter(submission__in=queryset)
        .select_related("field__section")
        .prefetch_related("field__options", "files")
        .order_by(
            "submission__form__name",
            "submission__form_id",
            "-submission__form_version__version_number",
            "field__section__order",
            "field__section_id",
            "field__order",
            "field_id",
        )
        .iterator(chunk_size=200)
    )
    columns = {}
    for answer in answers:
        if answer.value in (None, "", [], {}) and answer.field.field_type not in {
            "FILE",
            "DOCUMENT",
        }:
            continue
        if not report_answer_text(answer):
            continue
        label = answer.field.label
        key = answer_key(label)
        title = label.strip() or "Campo sin nombre"
        columns.setdefault(key, {"key": key, "title": title})
    return list(columns.values())


def overview_row(submission, columns):
    values = defaultdict(list)
    for answer in submission.summary_answers:
        value = report_answer_text(answer)
        if value:
            values[answer_key(answer.field.label)].append(value)
    return [
        timezone.localtime(submission.submitted_at).strftime("%Y-%m-%d %H:%M:%S"),
        *["\n".join(values[column["key"]]) for column in columns],
    ]


def document_row(submission, answer, attachment, path):
    name, document = identity(submission)
    return [
        name,
        document,
        answer.field.label,
        attachment.original_name,
        attachment.size,
        path,
    ]


def document_rows(submission):
    for _, answer, attachment, path in files_for([submission]):
        yield document_row(submission, answer, attachment, path)


def related_order(order, prefix):
    return tuple(
        f"-{prefix}{key[1:]}" if key.startswith("-") else f"{prefix}{key}"
        for key in REPORT_ORDER[order]
    )


def report_preview(model_admin, request):
    if request.method != "GET":
        return HttpResponseNotAllowed(["GET"])
    filters = report_filters(request)
    if filters is None:
        return HttpResponseBadRequest("Selecciona fechas, formulario y orden válidos.")
    sheet = request.GET.get("hoja", "respuestas")
    if sheet not in REPORT_SHEETS:
        sheet = "respuestas"
    queryset = scoped_queryset(model_admin, request, filters)
    response_count = queryset.count()
    forms_count = queryset.order_by().values("form_id").distinct().count()
    attachments = SubmissionFile.objects.filter(answer__submission__in=queryset)
    documents_count = attachments.count()
    latest_response = (
        queryset.order_by("-submitted_at").values_list("submitted_at", flat=True).first()
    )
    page_size = request.GET.get("por_pagina", "25")
    if page_size not in {"10", "25", "50"}:
        page_size = "25"
    if sheet == "respuestas":
        columns = report_columns(queryset)
        page = Paginator(queryset, int(page_size)).get_page(request.GET.get("pagina"))
        submissions = list(page.object_list)
        for submission in submissions:
            if not model_admin.has_view_permission(request, submission):
                raise PermissionDenied
        load_summary_data(submissions)
        rows = []
        for submission in submissions:
            name, document = identity(submission)
            values = overview_row(submission, columns)
            rows.append(
                {
                    "id": submission.pk,
                    "number": page.start_index() + len(rows),
                    "received": timezone.localtime(submission.submitted_at),
                    "form": submission.form.name,
                    "name": name,
                    "document": document,
                    "answers": [
                        {"key": column["key"], "value": value}
                        for column, value in zip(columns, values[1:])
                    ],
                }
            )
        sheet_headers = columns
    else:
        attachments = attachments.select_related(
            "answer__submission__form", "answer__field"
        ).order_by(
            *related_order(filters[2], "answer__submission__"),
            "answer__field__section__order",
            "answer__field__section_id",
            "answer__field__order",
            "answer__field_id",
            "pk",
        )
        page = Paginator(attachments, int(page_size)).get_page(request.GET.get("pagina"))
        files = list(page.object_list)
        submissions = {file.answer.submission_id: file.answer.submission for file in files}
        for submission in submissions.values():
            if not model_admin.has_view_permission(request, submission):
                raise PermissionDenied
        load_summary_data(list(submissions.values()))
        rows = []
        for file in files:
            submission = submissions[file.answer.submission_id]
            name, document = identity(submission)
            rows.append(
                {
                    "number": page.start_index() + len(rows),
                    "name": name,
                    "document": document,
                    "form": submission.form.name,
                    "field": file.answer.field.label,
                    "filename": file.original_name,
                    "type": file.original_name.rsplit(".", 1)[-1].upper()
                    if "." in file.original_name
                    else "Archivo",
                    "size": file.size,
                    "received": timezone.localtime(submission.submitted_at),
                    "url": file.get_absolute_url(),
                }
            )
        sheet_headers = []
    choices = list(
        model_admin.get_queryset(request)
        .order_by("form__name")
        .values_list("form_id", "form__name")
        .distinct()
    )
    download_query = request.GET.copy()
    for key in list(download_query):
        if key not in {"form", "desde", "hasta", "orden"}:
            download_query.pop(key, None)
    page_query = download_query.copy()
    page_query["hoja"] = sheet
    page_query["por_pagina"] = page_size
    tab_query = download_query.copy()
    tab_query["por_pagina"] = page_size
    selected_form_name = next(
        (name for pk, name in choices if pk == filters[1]), "Todos los formularios"
    )
    start, end = filters[0]
    if start and end:
        period_label = f"{start:%d/%m/%Y} – {end:%d/%m/%Y}"
    elif start:
        period_label = f"Desde {start:%d/%m/%Y}"
    elif end:
        period_label = f"Hasta {end:%d/%m/%Y}"
    else:
        period_label = "Todos los períodos"
    return render(
        request,
        "admin/submissions/reports.html",
        {
            **model_admin.admin_site.each_context(request),
            "opts": model_admin.model._meta,
            "title": "Reportes de respuestas",
            "form_choices": [
                {"id": str(pk), "name": name, "selected": pk == filters[1]} for pk, name in choices
            ],
            "selected_form": str(filters[1] or ""),
            "selected_order": filters[2],
            "selected_sheet": sheet,
            "sheet_headers": sheet_headers,
            "empty_colspan": len(sheet_headers) + 6 if sheet == "respuestas" else 10,
            "rows": rows,
            "page": page,
            "page_size": page_size,
            "response_count": response_count,
            "forms_count": forms_count,
            "documents_count": documents_count,
            "latest_response": latest_response,
            "selected_form_name": selected_form_name,
            "period_label": period_label,
            "download_query": download_query.urlencode(),
            "page_query": page_query.urlencode(),
            "tab_query": tab_query.urlencode(),
            "reports_url": reverse("admin:submissions_submission_reports"),
            "responses_url": reverse("admin:submissions_submission_changelist"),
        },
    )


def report_excel(model_admin, request):
    if request.method != "GET":
        return HttpResponseNotAllowed(["GET"])
    filters = report_filters(request)
    if filters is None:
        return HttpResponseBadRequest("Selecciona fechas, formulario y orden válidos.")
    queryset = scoped_queryset(model_admin, request, filters)
    columns = report_columns(queryset)
    submissions = prepared_submissions(model_admin, request, queryset)
    workbook = Workbook(write_only=True)
    overview = workbook.create_sheet("Respuestas")
    documents = workbook.create_sheet("Documentos")

    def header(sheet, labels):
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 34
        cells = []
        for label in labels:
            cell = WriteOnlyCell(sheet, value=safe_cell(label))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24324A")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cells.append(cell)
        sheet.append(cells)

    for index, width in enumerate([22, *([32] * len(columns))], 1):
        overview.column_dimensions[get_column_letter(index)].width = width
    for index, width in enumerate([27, 24, 30, 36, 18, 72], 1):
        documents.column_dimensions[get_column_letter(index)].width = width
    header(overview, [*REPORT_SHEETS["respuestas"], *(column["title"] for column in columns)])
    header(documents, REPORT_SHEETS["documentos"])
    for submission in submissions:
        overview.append([safe_cell(value) for value in overview_row(submission, columns)])
        for row in document_rows(submission):
            documents.append([safe_cell(value) for value in row])
    output = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024)
    workbook.save(output)
    output.seek(0)
    response = FileResponse(
        output,
        as_attachment=True,
        filename="reporte-respuestas.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["X-Content-Type-Options"] = "nosniff"
    ReportDownload.objects.create(actor=request.user, kind=ReportDownload.Kind.EXCEL)
    return response


def report_documents(model_admin, request, object_id=None):
    if request.method != "GET":
        return HttpResponseNotAllowed(["GET"])
    if not model_admin.has_view_permission(request):
        raise PermissionDenied
    if object_id:
        submission = get_object_or_404(model_admin.get_queryset(request), pk=object_id)
        if not model_admin.has_view_permission(request, submission):
            raise PermissionDenied
        submissions = [submission]
        load_summary_data(submissions)
    else:
        filters = report_filters(request)
        if filters is None:
            return HttpResponseBadRequest("Selecciona fechas, formulario y orden válidos.")
        submissions = scoped_submissions(model_admin, request, filters)
    response = StreamingHttpResponse(zip_chunks(submissions), content_type="application/zip")
    response["Content-Disposition"] = content_disposition_header(True, "documentos-respuestas.zip")
    response["X-Content-Type-Options"] = "nosniff"
    ReportDownload.objects.create(actor=request.user, kind=ReportDownload.Kind.DOCUMENTS)
    return response
