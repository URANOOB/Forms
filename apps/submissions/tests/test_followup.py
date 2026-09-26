import io
import uuid
from importlib import import_module
from unittest.mock import patch

from django.apps import apps
from django.contrib.auth.models import Group, Permission
from django.core.files.base import ContentFile
from django.core.files.storage import InMemoryStorage
from django.core.management import call_command
from django.db import connection
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.forms.models import FormField, FormSection, FormVersion
from apps.forms.publication import publish_form
from apps.submissions.admin_views import fingerprint
from apps.submissions.models import (
    PendingFileDeletion,
    Submission,
    SubmissionAnswer,
    SubmissionFile,
)
from apps.submissions.reports import answer_key, overview_row, report_columns, safe_cell
from apps.submissions.runtime import save_response
from apps.submissions.summary import load_summary_data
from apps.submissions.tests.test_public import fixture


class OwnershipAndTrashTests(TestCase):
    def setUp(self):
        self.form, self.version, self.name, *_ = fixture()
        self.owner = self.form.created_by
        self.other = User.objects.create_user(username="other", is_staff=True)
        self.admin = User.objects.create_user(username="admin", is_staff=True, is_superuser=True)
        for user in (self.owner, self.other):
            user.user_permissions.add(
                *Permission.objects.filter(
                    codename__in=[
                        "view_submission",
                        "change_submission",
                        "delete_submission",
                    ]
                )
            )
        self.submission = Submission.objects.create(
            form=self.form, form_version=self.version, idempotency_key=uuid.uuid4()
        )
        self.answer = SubmissionAnswer.objects.create(
            submission=self.submission, field=self.name, value="Original"
        )
        self.client.force_login(self.owner)

    def url(self, operation):
        return reverse(f"admin:submissions_submission_{operation}", args=[self.submission.pk])

    def trash(self):
        return self.client.post(self.url("remove"), {"confirm_delete": "yes"})

    def test_owner_rule_blocks_all_mutations_and_hides_edit_controls(self):
        self.client.post(self.url("review"), {"status": "UNDER_REVIEW", "revision": 0})
        self.submission.refresh_from_db()
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url("edit")).status_code, 403)
        for operation, data in [
            ("edit", {"revision": fingerprint(self.submission), "answer_nombre": "Changed"}),
            ("note", {"content": "Changed"}),
            ("attention", {"attention": "INCOMPLETE", "revision": self.submission.review_revision}),
            ("remove", {"confirm_delete": "yes"}),
        ]:
            with self.subTest(operation=operation):
                self.assertEqual(self.client.post(self.url(operation), data).status_code, 403)
        detail = self.client.get(self.url("detail"))
        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.context["can_edit"])
        self.answer.refresh_from_db()
        self.assertEqual(self.answer.value, "Original")
        for user in (self.owner, self.admin):
            self.client.force_login(user)
            self.submission.refresh_from_db()
            self.assertEqual(
                self.client.post(
                    self.url("edit"),
                    {"revision": fingerprint(self.submission), "answer_nombre": "Original"},
                ).status_code,
                302,
            )

    def test_trash_preserves_files_and_restore_recovers_data_and_history(self):
        field = SubmissionFile._meta.get_field("file")
        original_storage = field.storage
        storage = InMemoryStorage()
        field.storage = storage
        self.addCleanup(setattr, field, "storage", original_storage)
        attachment = SubmissionFile.objects.create(
            answer=self.answer,
            file=ContentFile(b"private", name="test.txt"),
            original_name="test.txt",
            size=7,
        )
        stale_revision = fingerprint(self.submission)
        self.assertEqual(self.trash().status_code, 302)
        self.assertFalse(Submission.objects.exists())
        self.assertTrue(Submission.all_objects.exists())
        self.assertTrue(storage.exists(attachment.file.name))
        self.assertTrue(self.submission.answers.exists())
        self.assertEqual(self.submission.activity.get().event_type, "trashed")
        for operation in ("detail", "edit", "download", "documents", "panel"):
            self.assertEqual(self.client.get(self.url(operation)).status_code, 404)
        self.assertEqual(self.client.get(attachment.get_absolute_url()).status_code, 404)
        reports = self.client.get(reverse("admin:submissions_submission_reports"))
        self.assertEqual(reports.context["response_count"], 0)
        self.assertEqual(reports.context["documents_count"], 0)
        board = self.client.get(reverse("admin:submissions_submission_changelist"))
        self.assertEqual(board.context["cl"].result_count, 0)
        trash = self.client.get(reverse("admin:submissions_submission_trash"))
        self.assertContains(trash, "Restaurar")
        self.assertNotContains(trash, self.url("purge"))
        self.assertEqual(self.client.post(self.url("restore")).status_code, 302)
        self.submission.refresh_from_db()
        self.assertIsNone(self.submission.deleted_at)
        self.assertNotEqual(fingerprint(self.submission), stale_revision)
        self.assertEqual(self.answer.value, "Original")
        response = self.client.get(attachment.get_absolute_url())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(response.streaming_content), b"private")
        self.assertTrue(self.submission.activity.filter(event_type="restored").exists())

    def test_only_admin_can_purge_and_only_after_trashing_and_confirmation(self):
        self.assertEqual(
            self.client.post(
                self.url("purge"), {"confirm_purge": str(self.submission.pk)}
            ).status_code,
            403,
        )
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self.url("purge")).status_code, 404)
        self.trash()
        self.assertEqual(self.client.get(self.url("purge")).status_code, 200)
        self.client.post(self.url("purge"), {"confirm_purge": "wrong"})
        self.assertTrue(Submission.all_objects.exists())
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                self.url("purge"), {"confirm_purge": str(self.submission.pk)}
            )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Submission.all_objects.exists())
        self.assertFalse(SubmissionAnswer.objects.exists())

    def test_purge_deletes_stored_files_after_commit(self):
        field = SubmissionFile._meta.get_field("file")
        storage = InMemoryStorage()
        previous = field.storage
        field.storage = storage
        self.addCleanup(setattr, field, "storage", previous)
        attachment = SubmissionFile.objects.create(
            answer=self.answer,
            file=ContentFile(b"private", name="test.txt"),
            original_name="test.txt",
            size=7,
        )
        self.trash()
        self.client.force_login(self.admin)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(self.url("purge"), {"confirm_purge": str(self.submission.pk)})
            self.assertTrue(storage.exists(attachment.file.name))
        self.assertFalse(storage.exists(attachment.file.name))

    def test_storage_failure_leaves_durable_purge_retry(self):
        field = SubmissionFile._meta.get_field("file")
        storage = InMemoryStorage()
        previous = field.storage
        field.storage = storage
        self.addCleanup(setattr, field, "storage", previous)
        attachment = SubmissionFile.objects.create(
            answer=self.answer,
            file=ContentFile(b"private", name="test.txt"),
            original_name="test.txt",
            size=7,
        )
        self.trash()
        self.client.force_login(self.admin)
        with patch.object(storage, "delete", side_effect=OSError("offline")):
            with self.assertLogs("apps.submissions.file_cleanup", level="ERROR"):
                with self.captureOnCommitCallbacks(execute=True):
                    self.client.post(self.url("purge"), {"confirm_purge": str(self.submission.pk)})
        self.assertFalse(Submission.all_objects.exists())
        self.assertEqual(PendingFileDeletion.objects.get().name, attachment.file.name)
        call_command("retry_file_deletions", stdout=io.StringIO())
        self.assertFalse(PendingFileDeletion.objects.exists())
        self.assertFalse(storage.exists(attachment.file.name))

    def test_trash_and_restore_require_permissions_and_csrf(self):
        secure_client = Client(enforce_csrf_checks=True)
        secure_client.force_login(self.owner)
        self.assertEqual(
            secure_client.post(self.url("remove"), {"confirm_delete": "yes"}).status_code, 403
        )
        self.trash()
        self.assertEqual(secure_client.post(self.url("restore")).status_code, 403)
        self.other.user_permissions.clear()
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(self.url("restore")).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("admin:submissions_submission_trash")).status_code, 403
        )

    def test_nonce_retries_do_not_resurrect_trashed_response(self):
        self.form = publish_form(self.form.pk)
        self.trash()
        retried = save_response(
            self.form.pk,
            self.version.pk,
            self.submission.idempotency_key,
            {self.name.pk: "Changed"},
        )
        self.assertEqual(retried.pk, self.submission.pk)
        self.assertIsNotNone(retried.deleted_at)
        self.assertEqual(Submission.all_objects.count(), 1)


class ReportIdentityTests(TestCase):
    def setUp(self):
        self.form, self.version, self.name, *_ = fixture()
        user = self.form.created_by
        user.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(user)

    def response(self, form, version, field, text):
        submission = Submission.objects.create(
            form=form, form_version=version, idempotency_key=uuid.uuid4()
        )
        SubmissionAnswer.objects.create(submission=submission, field=field, value=text)
        return submission

    def test_duplicate_labels_within_and_across_forms_have_distinct_columns(self):
        first = self.response(self.form, self.version, self.name, "Alice")
        duplicate = FormField.objects.create(
            form_version=self.version,
            section=self.name.section,
            stable_key="contact_name",
            label=self.name.label,
            field_type="SHORT_TEXT",
            order=8,
        )
        SubmissionAnswer.objects.create(submission=first, field=duplicate, value="Bob")
        form, version, field, *_ = fixture("second")
        second = self.response(form, version, field, "Carol")
        columns = report_columns(Submission.objects.all())
        self.assertEqual(len(columns), 3)
        self.assertEqual(len({column["title"] for column in columns}), 3)
        load_summary_data([first, second])
        for submission, expected in [(first, {"Alice", "Bob"}), (second, {"Carol"})]:
            values = overview_row(submission, columns)[1:]
            self.assertEqual(set(filter(None, values)), expected)
        preview = self.client.get(reverse("admin:submissions_submission_reports"))
        self.assertEqual(len(preview.context["sheet_headers"]), 3)

    def test_renamed_field_across_versions_retains_one_identity(self):
        first = self.response(self.form, self.version, self.name, "Old")
        publish_form(self.form.pk)
        version = FormVersion.objects.create(form=self.form, version_number=2)
        section = FormSection.objects.create(form_version=version, title="New")
        field = FormField.objects.create(
            form_version=version,
            section=section,
            stable_key=self.name.stable_key,
            label="New label",
            field_type="SHORT_TEXT",
        )
        second = self.response(self.form, version, field, "New")
        columns = report_columns(Submission.objects.all())
        self.assertEqual(columns[0]["title"], "New label")
        self.assertEqual(len(columns), 1)
        load_summary_data([first, second])
        self.assertEqual(overview_row(first, columns)[1:], ["Old"])
        self.assertEqual(overview_row(second, columns)[1:], ["New"])

    def test_excel_preserves_long_values_and_continuation_chunks_are_not_formulas(self):
        value = "A" * 32000 + "=1+1" + "\x01" * 10000
        submission = self.response(self.form, self.version, self.name, value)
        response = self.client.get(reverse("admin:submissions_submission_report_excel"))
        workbook = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        self.addCleanup(workbook.close)
        self.assertIn("Textos extensos", workbook["Respuestas"]["B2"].value)
        parts = list(workbook["Textos extensos"].iter_rows(min_row=2))
        self.assertEqual("".join(row[4].value for row in parts), safe_cell(value))
        for row in parts:
            self.assertEqual(row[1].value, str(submission.pk))
            self.assertEqual(row[2].value, answer_key(self.form.pk, self.name.stable_key))
            self.assertEqual(row[4].data_type, "s")
            self.assertLessEqual(len(row[4].value), 32767)


class OperatorMigrationTests(TestCase):
    def test_rename_and_merge_preserve_users_and_permissions(self):
        migration = import_module("apps.accounts.migrations.0005_rename_operator")
        old = Group.objects.create(name="Visor")
        user = User.objects.create_user(username="operator")
        user.groups.add(old)
        permission = Permission.objects.get(codename="change_submission")
        old.permissions.add(permission)
        with connection.schema_editor(atomic=False) as editor:
            migration.rename_operator(apps, editor)
            self.assertTrue(user.groups.filter(name="Operador").exists())
            self.assertTrue(user.has_perm("submissions.change_submission"))
            another_old = Group.objects.create(name="Visor")
            second = User.objects.create_user(username="second")
            second.groups.add(another_old)
            migration.rename_operator(apps, editor)
        self.assertTrue(second.groups.filter(name="Operador").exists())
        self.assertFalse(Group.objects.filter(name="Visor").exists())
