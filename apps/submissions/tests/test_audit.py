"""Regression coverage for data integrity, uploads and private exports."""

import io
import tempfile
import uuid
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.core.files.storage import FileSystemStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import SimpleTestCase, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.forms.conditions import FormSchema
from apps.forms.models import ConditionalRule, FormField
from apps.forms.public_fields import MAX_SAFE_INTEGER, public_field
from apps.forms.publication import publish_form
from apps.forms.question_fields import AttachmentField
from apps.submissions.models import Submission, SubmissionAnswer, SubmissionFile
from apps.submissions.reports import safe_cell
from apps.submissions.runtime import PublicResponseForm, new_token, save_response
from apps.submissions.tests.test_public import fixture


class InputAuditTests(SimpleTestCase):
    def number(self):
        return public_field(
            SimpleNamespace(
                field_type="NUMBER", validation={}, label="Cantidad", help_text="", placeholder=""
            )
        )

    def test_numbers_are_exact_integers_and_unsafe_values_are_rejected(self):
        field = self.number()
        for raw in ("0", "000123", str(MAX_SAFE_INTEGER)):
            with self.subTest(raw=raw):
                value = field.clean(raw)
                self.assertIs(type(value), int)
                self.assertEqual(value, int(raw))
        for raw in ("9007199254740993", "9" * 5000, "-1", "1.5", "1e2", "NaN"):
            with self.subTest(raw=raw), self.assertRaises(ValidationError):
                field.clean(raw)

    def test_excel_controls_and_formulas_round_trip_as_text(self):
        workbook = Workbook()
        values = ["Ana\x0bPérez", "\ufffe\uffff", "=1+1", "\x01=1+1", "línea\nsegunda\tcolumna"]
        for value in values:
            workbook.active.append([safe_cell(value)])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        loaded = load_workbook(output)
        self.assertEqual(loaded.active["A1"].value, "Ana\\u000bPérez")
        self.assertEqual(loaded.active["A2"].value, "\\ufffe\\uffff")
        self.assertTrue(all(loaded.active.cell(i, 1).data_type == "s" for i in range(1, 6)))
        self.assertEqual(loaded.active["A5"].value, values[-1])
        loaded.close()

    def attachment(self):
        return AttachmentField(SimpleNamespace(configuration={}), required=False)

    def office_zip(self, entries, name="document.docx"):
        output = io.BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            for key, value in entries.items():
                archive.writestr(key, value)
        return SimpleUploadedFile(name, output.getvalue())

    def test_office_rejects_renamed_zip_truncation_and_macros(self):
        valid_parts = {
            "[Content_Types].xml": "<Types/>",
            "_rels/.rels": "<Relationships/>",
            "word/document.xml": "<document/>",
        }
        invalid = [
            self.office_zip({"program.exe": "MZ"}),
            SimpleUploadedFile("broken.docx", b"PK\x03\x04broken"),
            self.office_zip({**valid_parts, "word/vbaProject.bin": "macro"}),
            self.office_zip(valid_parts, "wrong.xlsx"),
        ]
        for upload in invalid:
            with self.subTest(name=upload.name), self.assertRaises(ValidationError):
                self.attachment().clean(upload)

    def test_real_excel_is_accepted_and_stream_is_rewound(self):
        output = io.BytesIO()
        workbook = Workbook()
        workbook.active.append(["Documento", "00123"])
        workbook.save(output)
        upload = SimpleUploadedFile("document.xlsx", output.getvalue())
        self.assertEqual(self.attachment().clean(upload), [upload])
        self.assertEqual(upload.tell(), 0)


class ResponseAuditTests(TestCase):
    def setUp(self):
        self.form, self.version, self.name, self.choice, self.email = fixture()
        self.user = self.form.created_by
        self.user.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(self.user)

    def create_response(self, value="Ana", **kwargs):
        submission = Submission.objects.create(
            form=self.form, form_version=self.version, idempotency_key=uuid.uuid4(), **kwargs
        )
        SubmissionAnswer.objects.create(submission=submission, field=self.name, value=value)
        return submission

    def test_excel_endpoint_handles_control_characters_and_enforces_permission(self):
        self.create_response("Ana\x0bPérez")
        url = reverse("admin:submissions_submission_report_excel")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(workbook["Respuestas"]["B2"].value, "Ana\\u000bPérez")
        workbook.close()
        self.user.user_permissions.clear()
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_reports_include_both_bogota_date_boundaries(self):
        for hour, minute, day in [(4, 59, 25), (5, 0, 25), (4, 59, 26), (5, 0, 26)]:
            self.create_response(
                submitted_at=datetime(2026, 9, day, hour, minute, tzinfo=timezone.utc)
            )
        response = self.client.get(
            reverse("admin:submissions_submission_reports"),
            {"desde": "2026-09-25", "hasta": "2026-09-25"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["response_count"], 2)
        self.assertEqual(len(response.context["rows"]), 2)

    def test_invalid_report_filters_fail_without_server_error(self):
        url = reverse("admin:submissions_submission_report_excel")
        for params in (
            {"form": "invalid"},
            {"desde": "2026-02-30"},
            {"desde": "2026-09-26", "hasta": "2026-09-25"},
            {"orden": "invalid"},
        ):
            with self.subTest(params=params):
                self.assertEqual(self.client.get(url, params).status_code, 400)

    def test_numbers_cannot_be_silently_rounded_by_public_submission(self):
        numeric = FormField.objects.create(
            form_version=self.version,
            section=self.name.section,
            label="Cantidad",
            stable_key="cantidad",
            field_type="NUMBER",
            order=4,
        )
        self.form = publish_form(self.form.pk)
        payload = {
            "submission_token": new_token(self.form),
            "answer_nombre": "Ana",
            "answer_contactar": "no",
            "answer_cantidad": "9007199254740993",
        }
        self.assertEqual(self.client.post(self.form.get_absolute_url(), payload).status_code, 422)
        self.assertFalse(Submission.objects.exists())
        payload["answer_cantidad"] = str(MAX_SAFE_INTEGER)
        self.assertEqual(self.client.post(self.form.get_absolute_url(), payload).status_code, 302)
        self.assertEqual(SubmissionAnswer.objects.get(field=numeric).value, MAX_SAFE_INTEGER)

    def test_historical_large_numeric_threshold_does_not_break_public_form(self):
        numeric = FormField.objects.create(
            form_version=self.version,
            section=self.name.section,
            label="Cantidad",
            stable_key="cantidad",
            field_type="NUMBER",
            order=4,
        )
        rule = ConditionalRule.objects.create(
            form_version=self.version,
            source_field=numeric,
            target_field=self.name,
            operator="GREATER_THAN",
            expected_value=10,
            action="REQUIRE",
        )
        self.form = publish_form(self.form.pk)
        # Simulate a threshold accepted and published by the previous runtime.
        ConditionalRule.objects.filter(pk=rule.pk).update(expected_value=MAX_SAFE_INTEGER + 2)
        self.assertEqual(self.client.get(self.form.get_absolute_url()).status_code, 200)

    def test_answer_inserts_are_batched_and_retry_does_not_duplicate(self):
        fields = [
            FormField.objects.create(
                form_version=self.version,
                section=self.name.section,
                label=f"Dato {i}",
                stable_key=f"dato_{i}",
                field_type="SHORT_TEXT",
                order=i + 4,
            )
            for i in range(20)
        ]
        self.form = publish_form(self.form.pk)
        nonce = uuid.uuid4()
        answers = {field.pk: f"valor {index}" for index, field in enumerate(fields)}
        with CaptureQueriesContext(connection) as queries:
            submission = save_response(self.form.pk, self.version.pk, nonce, answers)
        inserts = [
            query["sql"]
            for query in queries
            if query["sql"].startswith('INSERT INTO "submissions_submissionanswer"')
        ]
        self.assertEqual(len(inserts), 1)
        self.assertEqual(submission.answers.count(), 20)
        self.assertEqual(
            save_response(self.form.pk, self.version.pk, nonce, answers).pk, submission.pk
        )
        self.assertEqual(Submission.objects.count(), 1)

    def test_private_download_headers_and_permission(self):
        submission = self.create_response()
        with tempfile.TemporaryDirectory() as directory:
            storage = FileSystemStorage(location=directory)
            name = storage.save("private.txt", SimpleUploadedFile("private.txt", b"private"))
            attachment = SubmissionFile.objects.create(
                answer=submission.answers.get(), file=name, original_name="private.txt", size=7
            )
            with patch.object(SubmissionFile._meta.get_field("file"), "storage", storage):
                response = self.client.get(attachment.get_absolute_url())
                self.assertEqual(response.status_code, 200)
                self.assertEqual(b"".join(response.streaming_content), b"private")
                self.assertTrue(response["Content-Disposition"].startswith("attachment"))
                self.assertIn("no-store", response["Cache-Control"])
                self.assertEqual(response["Content-Security-Policy"], "sandbox")
                self.user.user_permissions.clear()
                self.assertEqual(self.client.get(attachment.get_absolute_url()).status_code, 403)

    @override_settings(SUBMISSION_MAX_BYTES=4_000_000)
    def test_aggregate_upload_limit_rejects_individually_valid_files(self):
        FormField.objects.create(
            form_version=self.version,
            section=self.name.section,
            label="Archivos",
            stable_key="archivos",
            field_type="FILE",
            order=4,
            configuration={"max_files": 2, "extensions": ["txt"]},
        )
        schema = FormSchema(self.version)
        response = PublicResponseForm(
            schema,
            data={"answer_nombre": "Ana", "answer_contactar": "no"},
            files={
                "answer_archivos": [
                    SimpleUploadedFile(f"{i}.txt", b"a" * 2_000_000) for i in range(2)
                ]
            },
        )
        self.assertFalse(response.is_valid())
        self.assertIn("respuesta completa", str(response.non_field_errors()))
        self.assertNotIn("answer_archivos", response.errors)
        self.assertEqual(schema.browser_spec()["max_submission_bytes"], 4_000_000)
